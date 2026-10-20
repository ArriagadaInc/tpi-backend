"""Construye las tablas de mortalidad del simulador desde el Excel oficial de la CMF.

Genera ``app/simulador_pension/datos/tablas_mortalidad.npz`` con, para cada tipo (vejez, sobrev,
inval) y sexo (H, M):

    <tipo>_<sexo>_q    (111,)     q(edad) base 2020
    <tipo>_<sexo>_aa   (111, 16)  factores de mejoramiento anual AA(edad, año 2021..2036)

Hojas y columnas del Excel (índice base 0 de la columna «Edad»; q va en la siguiente y los 16 AA en
las que siguen). Se lee solo la serie 2020 de cada hoja:

    vejez   Vejez-Hombres (CB-2020) / Vejez-Mujeres (RV-2020)            columna 12
    sobrev  Sobrevivencia-Hombres (CB-2020) / Sobrevivencia-Mujeres (B-2020)   columna 8
    inval   Invalidez-Hombres / Invalidez-Mujeres (MI-2020)              columna 8

Antes de leer, el script comprueba los encabezados de cada bloque. Si la CMF reordena el archivo o
agrega series, falla con un mensaje claro en lugar de leer columnas equivocadas.

Uso (desde la raíz del repositorio):

    python scripts/tablas_mortalidad/construir_tablas.py                 # Excel del repo -> .npz
    python scripts/tablas_mortalidad/construir_tablas.py --verificar     # no escribe: compara
    python scripts/tablas_mortalidad/construir_tablas.py <excel.xlsx>    # otro Excel
    python scripts/tablas_mortalidad/construir_tablas.py --salida <ruta.npz>

Requiere ``openpyxl`` (ya es dependencia de ejecución). Procedimiento completo y procedencia del
Excel en ``scripts/tablas_mortalidad/README.md`` y ``docs/SIMULADOR_USO.md`` (§8.2).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl

EDAD_MAX = 110
N_AA = 16  # años de mejoramiento 2021..2036
PRIMER_ANIO_AA = 2021
SERIE = 2020

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
EXCEL_POR_DEFECTO = AQUI / "tablas-mortalidad-cmf.xlsx"
PROCEDENCIA = AQUI / "procedencia.json"
NPZ_POR_DEFECTO = RAIZ / "app" / "simulador_pension" / "datos" / "tablas_mortalidad.npz"

# tipo -> (hoja de hombres, hoja de mujeres, columna base 0 de «Edad» de la serie 2020)
HOJAS: dict[str, tuple[str, str, int]] = {
    "vejez": ("Vejez-Hombres", "Vejez-Mujeres", 12),
    "sobrev": ("Sobrevivencia-Hombres", "Sobrevivencia-Mujeres", 8),
    "inval": ("Invalidez-Hombres", "Invalidez-Mujeres", 8),
}


class ExcelInesperadoError(ValueError):
    """El Excel no tiene la estructura que el script sabe leer."""


def sha256_archivo(ruta: Path) -> str:
    """Huella SHA-256 del archivo, en hexadecimal."""
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def _celda(ws: Any, fila: int, columna0: int) -> Any:
    """Valor de una celda con fila base 1 y columna base 0."""
    return ws.cell(row=fila, column=columna0 + 1).value


def _anio_de_encabezado(valor: Any) -> int | None:
    """Año de un encabezado de AA. La CMF rotula el último como «2036*» (texto con asterisco)."""
    try:
        return int(str(valor).rstrip("*").strip())
    except ValueError:
        return None


def validar_encabezado(ws: Any, columna: int) -> None:
    """Comprueba que el bloque de la serie 2020 está donde el script lo espera."""
    titulo = _celda(ws, 1, columna)
    if not isinstance(titulo, str) or str(SERIE) not in titulo:
        raise ExcelInesperadoError(
            f"hoja {ws.title!r}: la columna {columna} no abre la serie {SERIE} (título: {titulo!r})"
        )
    esperado = {columna: "Edad", columna + 1: f"qx {SERIE}"}
    for col, texto in esperado.items():
        valor = _celda(ws, 2, col)
        if valor != texto:
            raise ExcelInesperadoError(
                f"hoja {ws.title!r}: se esperaba {texto!r} en la fila 2, columna {col}; hay {valor!r}"
            )
    anios = [_anio_de_encabezado(_celda(ws, 3, columna + 2 + k)) for k in range(N_AA)]
    if anios != list(range(PRIMER_ANIO_AA, PRIMER_ANIO_AA + N_AA)):
        raise ExcelInesperadoError(
            f"hoja {ws.title!r}: los años de mejoramiento no son {PRIMER_ANIO_AA}..."
            f"{PRIMER_ANIO_AA + N_AA - 1} (hay {anios!r})"
        )


def leer_hoja(ws: Any, columna: int) -> tuple[np.ndarray, np.ndarray]:
    """(q[edad], aa[edad, año]) de una hoja. Las edades sin dato quedan como NaN en q."""
    validar_encabezado(ws, columna)
    q = np.full(EDAD_MAX + 1, np.nan)
    aa = np.zeros((EDAD_MAX + 1, N_AA))
    for fila in ws.iter_rows(min_row=4, values_only=True):
        edad, valor_q = fila[columna], fila[columna + 1]
        if isinstance(edad, (int, float)) and isinstance(valor_q, (int, float)):
            e = int(edad)
            q[e] = valor_q
            aa[e, :] = [0.0 if v is None else v for v in fila[columna + 2 : columna + 2 + N_AA]]
    return q, aa


def construir(ruta_xlsx: Path) -> dict[str, np.ndarray]:
    """Lee el Excel y devuelve los 12 arreglos que guarda el ``.npz``."""
    libro = openpyxl.load_workbook(ruta_xlsx, read_only=True, data_only=True)
    salida: dict[str, np.ndarray] = {}
    for tipo, (hoja_h, hoja_m, columna) in HOJAS.items():
        for sexo, hoja in (("H", hoja_h), ("M", hoja_m)):
            if hoja not in libro.sheetnames:
                raise ExcelInesperadoError(f"el Excel no tiene la hoja {hoja!r}")
            q, aa = leer_hoja(libro[hoja], columna)
            salida[f"{tipo}_{sexo}_q"] = q
            salida[f"{tipo}_{sexo}_aa"] = aa
    return salida


def diferencias(nuevo: dict[str, np.ndarray], ruta_npz: Path) -> list[str]:
    """Lista legible de las diferencias entre lo construido y un ``.npz`` existente."""
    actual = dict(np.load(ruta_npz))
    problemas: list[str] = []
    for clave in sorted(set(nuevo) | set(actual)):
        if clave not in nuevo or clave not in actual:
            problemas.append(
                f"{clave}: solo existe en {'el Excel' if clave in nuevo else 'el .npz'}"
            )
        elif not np.array_equal(nuevo[clave], actual[clave], equal_nan=True):
            problemas.append(f"{clave}: los valores difieren")
    return problemas


def comprobar_procedencia(ruta_xlsx: Path) -> str | None:
    """Aviso si el Excel no es el registrado en ``procedencia.json`` (None si coincide)."""
    if not PROCEDENCIA.is_file():
        return "no existe procedencia.json"
    registrada = json.loads(PROCEDENCIA.read_text(encoding="utf-8"))["sha256"]
    if sha256_archivo(ruta_xlsx) != registrada:
        return (
            "el Excel NO es el registrado en procedencia.json; si es una versión nueva de la CMF, "
            "actualiza procedencia.json (url, fecha, tamaño y sha256)"
        )
    return None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=(__doc__ or "").split("\n")[0])
    parser.add_argument("excel", nargs="?", type=Path, default=EXCEL_POR_DEFECTO)
    parser.add_argument("--salida", type=Path, default=NPZ_POR_DEFECTO)
    parser.add_argument(
        "--verificar",
        action="store_true",
        help="no escribe nada: comprueba que el .npz coincide con lo que sale del Excel",
    )
    args = parser.parse_args(argv)

    aviso = comprobar_procedencia(args.excel)
    if aviso:
        print(f"AVISO: {aviso}", file=sys.stderr)
    try:
        tablas = construir(args.excel)
    except ExcelInesperadoError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    if args.verificar:
        problemas = diferencias(tablas, args.salida)
        if problemas:
            print("El .npz NO coincide con el Excel:")
            for p in problemas:
                print(f"  - {p}")
            return 1
        print(f"OK: {args.salida.name} coincide con {args.excel.name} ({len(tablas)} arreglos).")
        return 0

    np.savez_compressed(args.salida, **tablas)
    print(f"escrito {args.salida} ({args.salida.stat().st_size} bytes, {len(tablas)} arreglos)")
    print("Recuerda: subir __version__ en app/simulador_pension/__init__.py y correr la regresión.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
