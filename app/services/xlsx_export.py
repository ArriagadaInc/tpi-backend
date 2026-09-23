"""Ephemeral, in-memory XLSX builder for the executive lead export (H3.3.6).

The builder serializes **only** the explicit ``EXPORT_COLUMNS`` allowlist. It never
iterates over row keys, never uses reflection and never discovers columns
dynamically. The workbook is assembled in memory and returned as raw bytes: no
temporary file, no persistence, nothing left behind on error.

Formula-injection defense: text values whose first character is ``=``, ``+``, ``-``
or ``@`` are neutralized with a leading single quote and the cell is forced to a
text number format. Numbers and dates are written as native cell types and are
never neutralized.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

EXPORT_SHEET_NAME = "Leads"
# Documented in docs/H3_3_6_EXPORT_COLUMNS.md. Count-first abort keeps openpyxl's
# in-memory footprint bounded far above the synthetic DEV volume.
EXPORT_MAX_ROWS = 10000

_TEXT = "text"
_NUMBER = "number"
_DATE = "date"

_FORMULA_TRIGGER_PREFIXES = ("=", "+", "-", "@")
_TEXT_FORMAT = "@"


@dataclass(frozen=True, slots=True)
class ExportColumn:
    """One allowlisted column: source field, stable header, cell kind and format."""

    field: str
    header: str
    kind: str
    number_format: str | None = None


# Explicit allowlist, order is canonical. Must stay in sync with
# docs/H3_3_6_EXPORT_COLUMNS.md and tests/unit/test_h3_3_6_requirements_matrix_gate.py.
EXPORT_COLUMNS: tuple[ExportColumn, ...] = (
    ExportColumn("id_lead", "ID Lead", _TEXT),
    ExportColumn("rut", "RUT", _TEXT),
    ExportColumn("nombre_completo", "Nombre", _TEXT),
    ExportColumn("email", "Email", _TEXT),
    ExportColumn("telefono", "Teléfono", _TEXT),
    ExportColumn("genero", "Género", _TEXT),
    ExportColumn("estado_civil", "Estado Civil", _TEXT),
    ExportColumn("afp", "AFP", _TEXT),
    ExportColumn("saldo_afp", "Saldo AFP (CLP)", _NUMBER, "#,##0"),
    ExportColumn("comentarios", "Comentarios", _TEXT),
    ExportColumn("estado_lead", "Estado", _TEXT),
    ExportColumn("created_at", "Fecha Ingreso", _DATE, "dd/mm/yyyy"),
    ExportColumn("id_asesor", "ID Asesor", _TEXT),
    ExportColumn("asesor_nombre", "Asesor", _TEXT),
)

_COLUMN_WIDTHS: dict[str, int] = {
    "id_lead": 38,
    "rut": 14,
    "nombre_completo": 30,
    "email": 30,
    "telefono": 18,
    "genero": 16,
    "estado_civil": 16,
    "afp": 16,
    "saldo_afp": 16,
    "comentarios": 40,
    "estado_lead": 16,
    "created_at": 14,
    "id_asesor": 38,
    "asesor_nombre": 24,
}


class ExportLimitExceededError(Exception):
    """Raised when the matching result count exceeds ``EXPORT_MAX_ROWS``."""

    def __init__(self, total: int, limit: int) -> None:
        self.total = total
        self.limit = limit
        super().__init__(f"Exportacion excede el limite de {limit} filas (coinciden {total})")


def sanitize_cell_text(value: Any) -> str:
    """Neutralize text that could be interpreted as a formula prefix.

    ``None`` becomes the empty string. Any value whose string form starts with
    ``=``, ``+``, ``-`` or ``@`` is prefixed with a single quote, the canonical
    spreadsheet text marker. The result is always a plain string, never a formula.
    """
    if value is None:
        return ""
    text = str(value)
    if text.startswith(_FORMULA_TRIGGER_PREFIXES):
        return "'" + text
    return text


def _coerce_number(value: Any) -> int | float | None:
    """Coerce a numeric DB value to a native spreadsheet number, or ``None``."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value
    if isinstance(value, Decimal):
        integral = value.to_integral_value()
        return int(integral) if value == integral else float(value)
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else number


def _coerce_date(value: Any) -> datetime | date | None:
    """Coerce a DB timestamp/date to a native spreadsheet date, or ``None``."""
    if isinstance(value, (datetime, date)):
        return value
    return None


def build_xlsx_workbook(rows: list[dict[str, Any]]) -> bytes:
    """Build the export workbook in memory and return its bytes.

    ``rows`` must be the full matching result set (already limit-checked by the
    caller). Extra keys in a row are ignored: only ``EXPORT_COLUMNS`` are read.
    """
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = EXPORT_SHEET_NAME

    header_font = Font(bold=True)
    for column_index, column in enumerate(EXPORT_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=column.header)
        cell.font = header_font
        if column.kind == _TEXT:
            cell.number_format = _TEXT_FORMAT

    worksheet.freeze_panes = "A2"
    last_column = get_column_letter(len(EXPORT_COLUMNS))
    worksheet.auto_filter.ref = f"A1:{last_column}{max(len(rows) + 1, 1)}"

    for column_index, column in enumerate(EXPORT_COLUMNS, start=1):
        width = _COLUMN_WIDTHS.get(column.field, 18)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(EXPORT_COLUMNS, start=1):
            value = row.get(column.field) if isinstance(row, dict) else None
            cell = worksheet.cell(row=row_index, column=column_index)
            if column.kind == _NUMBER:
                cell.value = _coerce_number(value)
                cell.number_format = column.number_format or "General"
            elif column.kind == _DATE:
                cell.value = _coerce_date(value)
                cell.number_format = column.number_format or "dd/mm/yyyy"
            else:
                cell.value = sanitize_cell_text(value)
                cell.number_format = _TEXT_FORMAT

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
