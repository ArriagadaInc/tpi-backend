"""Unit tests for the H3.3.6 XLSX export builder.

Verifies the explicit allowlist, cell types, formats, formula-injection
neutralization, Unicode, leading zeros, nulls and the empty-result case by
writing a workbook in memory and reading it back with openpyxl.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.services.xlsx_export import (
    EXPORT_COLUMNS,
    EXPORT_MAX_ROWS,
    EXPORT_SHEET_NAME,
    ExportLimitExceededError,
    build_xlsx_workbook,
    sanitize_cell_text,
)

EXPECTED_HEADERS = [column.header for column in EXPORT_COLUMNS]
_FIELD_TO_COLUMN = {column.field: index for index, column in enumerate(EXPORT_COLUMNS, start=1)}


def _base_row(**overrides: Any) -> dict[str, Any]:
    row: dict[str, Any] = {
        "id_lead": "11111111-1111-1111-1111-111111111111",
        "rut": "12.345.678-5",
        "nombre_completo": "Juan Pérez",
        "email": "juan@example.com",
        "telefono": "+56 9 1234 5678",
        "genero": "Masculino",
        "estado_civil": "Soltero",
        "afp": "Habitat",
        "saldo_afp": 5120000,
        "comentarios": "Comentario inicial",
        "estado_lead": "contactado",
        "created_at": datetime(2026, 9, 1, 10, 30),
        "id_asesor": "aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa",
        "asesor_nombre": "Asesor Uno",
    }
    row.update(overrides)
    return row


def _load_sheet(content: bytes):
    workbook = load_workbook(BytesIO(content))
    return workbook, workbook[EXPORT_SHEET_NAME]


def _cell(worksheet, row: int, field: str):
    return worksheet.cell(row=row, column=_FIELD_TO_COLUMN[field])


def test_build_returns_valid_xlsx_single_sheet() -> None:
    workbook, sheet = _load_sheet(build_xlsx_workbook([_base_row()]))
    assert workbook.sheetnames == [EXPORT_SHEET_NAME]
    assert sheet.title == EXPORT_SHEET_NAME


def test_headers_match_allowlist_exactly() -> None:
    _, sheet = _load_sheet(build_xlsx_workbook([]))
    headers = [sheet.cell(row=1, column=i).value for i in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPECTED_HEADERS


def test_freeze_panes_and_autofilter() -> None:
    _, sheet = _load_sheet(build_xlsx_workbook([_base_row()]))
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None
    assert sheet.auto_filter.ref.startswith("A1:")


def test_column_widths_configured() -> None:
    _, sheet = _load_sheet(build_xlsx_workbook([]))
    assert sheet.column_dimensions["A"].width is not None
    assert sheet.column_dimensions["B"].width is not None


def test_formula_trigger_prefixes_neutralized() -> None:
    row = _base_row(
        nombre_completo="=1+1",
        email="+cmd|' /C calc'!A0",
        telefono="-2+3",
        comentarios="@SUM(1,2)",
    )
    _, sheet = _load_sheet(build_xlsx_workbook([row]))
    assert _cell(sheet, 2, "nombre_completo").value == "'=1+1"
    assert _cell(sheet, 2, "email").value == "'+cmd|' /C calc'!A0"
    assert _cell(sheet, 2, "telefono").value == "'-2+3"
    assert _cell(sheet, 2, "comentarios").value == "'@SUM(1,2)"
    for field in ("nombre_completo", "email", "telefono", "comentarios"):
        assert _cell(sheet, 2, field).data_type == "s"


def test_sanitize_cell_text_helper() -> None:
    assert sanitize_cell_text("=1+1") == "'=1+1"
    assert sanitize_cell_text("+1") == "'+1"
    assert sanitize_cell_text("-1") == "'-1"
    assert sanitize_cell_text("@cmd") == "'@cmd"
    assert sanitize_cell_text("normal") == "normal"
    assert sanitize_cell_text(None) == ""
    assert sanitize_cell_text(123) == "123"


def test_numbers_native_clp_format() -> None:
    _, sheet = _load_sheet(build_xlsx_workbook([_base_row(saldo_afp=5120000)]))
    cell = _cell(sheet, 2, "saldo_afp")
    assert cell.data_type == "n"
    assert cell.value == 5120000
    assert cell.number_format == "#,##0"


def test_dates_native_date_cells() -> None:
    created = datetime(2026, 9, 1, 10, 30)
    _, sheet = _load_sheet(build_xlsx_workbook([_base_row(created_at=created)]))
    cell = _cell(sheet, 2, "created_at")
    assert cell.data_type == "d"
    assert cell.number_format == "dd/mm/yyyy"


def test_uuids_rut_phone_leading_zeros_as_text() -> None:
    row = _base_row(
        rut="00123456-7",
        telefono="+56 9 0123 4567",
        id_lead="00000000-0000-0000-0000-000000000001",
    )
    _, sheet = _load_sheet(build_xlsx_workbook([row]))
    assert _cell(sheet, 2, "rut").value == "00123456-7"
    assert _cell(sheet, 2, "rut").data_type == "s"
    assert _cell(sheet, 2, "telefono").value == "+56 9 0123 4567"
    assert _cell(sheet, 2, "telefono").data_type == "s"
    assert _cell(sheet, 2, "id_lead").value == "00000000-0000-0000-0000-000000000001"


def test_unicode_survives_roundtrip() -> None:
    row = _base_row(nombre_completo="Ñandú García", comentarios="Café – emoji 😀 y acentos áéíóú")
    _, sheet = _load_sheet(build_xlsx_workbook([row]))
    assert _cell(sheet, 2, "nombre_completo").value == "Ñandú García"
    assert _cell(sheet, 2, "comentarios").value == "Café – emoji 😀 y acentos áéíóú"


def test_null_values_become_empty_text() -> None:
    _, sheet = _load_sheet(build_xlsx_workbook([_base_row(email=None, id_asesor=None)]))
    assert _cell(sheet, 2, "email").value in ("", None)
    assert _cell(sheet, 2, "id_asesor").value in ("", None)


def test_empty_result_headers_only() -> None:
    _, sheet = _load_sheet(build_xlsx_workbook([]))
    assert sheet.max_row == 1
    headers = [sheet.cell(row=1, column=i).value for i in range(1, len(EXPORT_COLUMNS) + 1)]
    assert headers == EXPECTED_HEADERS


def test_no_formula_cells() -> None:
    _, sheet = _load_sheet(build_xlsx_workbook([_base_row(nombre_completo="=cmd")]))
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.data_type != "f", f"celda con fórmula en {cell.coordinate}"


def test_export_limit_error_carries_total_and_limit() -> None:
    error = ExportLimitExceededError(total=EXPORT_MAX_ROWS + 1, limit=EXPORT_MAX_ROWS)
    assert error.total == EXPORT_MAX_ROWS + 1
    assert error.limit == EXPORT_MAX_ROWS
    assert str(EXPORT_MAX_ROWS) in str(error)
