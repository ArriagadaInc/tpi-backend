"""Las tablas de mortalidad del simulador salen del Excel de la CMF versionado en el repositorio.

Protegen dos cosas: que el Excel del repo es el registrado en ``procedencia.json`` y que el
``.npz`` que usa el motor se reconstruye exactamente desde ese Excel. Si alguien cambia uno sin
el otro, el CI falla. Procedimiento de actualización en ``docs/SIMULADOR_USO.md`` (§8.2).
"""

from __future__ import annotations

import importlib.util
import json
from pathlib import Path
from types import ModuleType

import numpy as np
import openpyxl
import pytest

RAIZ = Path(__file__).resolve().parents[2]
CARPETA = RAIZ / "scripts" / "tablas_mortalidad"
NPZ = RAIZ / "app" / "simulador_pension" / "datos" / "tablas_mortalidad.npz"


def _cargar_script() -> ModuleType:
    spec = importlib.util.spec_from_file_location(
        "construir_tablas", CARPETA / "construir_tablas.py"
    )
    assert spec is not None and spec.loader is not None
    modulo = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modulo)
    return modulo


@pytest.fixture(scope="module")
def script() -> ModuleType:
    return _cargar_script()


@pytest.fixture(scope="module")
def tablas(script: ModuleType) -> dict[str, np.ndarray]:
    return script.construir(script.EXCEL_POR_DEFECTO)


def test_el_excel_del_repo_es_el_registrado(script: ModuleType) -> None:
    procedencia = json.loads((CARPETA / "procedencia.json").read_text(encoding="utf-8"))
    excel = CARPETA / procedencia["archivo"]
    assert excel == script.EXCEL_POR_DEFECTO
    assert excel.stat().st_size == procedencia["bytes"]
    assert script.sha256_archivo(excel) == procedencia["sha256"]
    assert script.comprobar_procedencia(excel) is None


def test_el_npz_del_motor_se_reconstruye_desde_el_excel(
    script: ModuleType, tablas: dict[str, np.ndarray]
) -> None:
    assert script.diferencias(tablas, NPZ) == []


def test_estructura_de_las_tablas(tablas: dict[str, np.ndarray]) -> None:
    esperadas = {
        f"{t}_{s}_{k}" for t in ("vejez", "sobrev", "inval") for s in "HM" for k in ("q", "aa")
    }
    assert set(tablas) == esperadas
    for clave, valores in tablas.items():
        if clave.endswith("_q"):
            assert valores.shape == (111,)
            assert valores[110] == 1.0  # edad terminal: se muere con certeza
            # Las edades que usa el simulador (a partir de 46 años) tienen dato y son probabilidades.
            usadas = valores[46:]
            assert not np.isnan(usadas).any()
            assert ((usadas > 0) & (usadas <= 1)).all()
        else:
            assert valores.shape == (111, 16)
            assert (valores >= 0).all()


def test_la_tabla_de_mujeres_pensionadas_parte_en_la_edad_20(tablas: dict[str, np.ndarray]) -> None:
    q = tablas["vejez_M_q"]
    assert np.isnan(q[:20]).all()
    assert not np.isnan(q[20:]).any()


def test_diferencias_detecta_un_cambio(
    script: ModuleType, tablas: dict[str, np.ndarray], tmp_path: Path
) -> None:
    alterado = {k: v.copy() for k, v in tablas.items()}
    alterado["vejez_H_q"][70] += 0.001
    ruta = tmp_path / "alterado.npz"
    np.savez_compressed(ruta, **alterado)
    assert script.diferencias(tablas, ruta) == ["vejez_H_q: los valores difieren"]


def test_un_excel_sin_las_hojas_esperadas_falla_con_mensaje_claro(
    script: ModuleType, tmp_path: Path
) -> None:
    libro = openpyxl.Workbook()
    libro.active.title = "Otra"
    ruta = tmp_path / "otro.xlsx"
    libro.save(ruta)
    with pytest.raises(script.ExcelInesperadoError, match="no tiene la hoja"):
        script.construir(ruta)


def test_un_excel_con_encabezados_distintos_falla_con_mensaje_claro(
    script: ModuleType, tmp_path: Path
) -> None:
    libro = openpyxl.Workbook()
    primera = True
    for hoja_h, hoja_m, _ in script.HOJAS.values():
        for nombre in (hoja_h, hoja_m):
            hoja = libro.active if primera else libro.create_sheet()
            hoja.title = nombre
            primera = False
    ruta = tmp_path / "encabezados.xlsx"
    libro.save(ruta)
    with pytest.raises(script.ExcelInesperadoError, match="no abre la serie 2020"):
        script.construir(ruta)


def test_main_verificar_devuelve_cero_si_coincide_y_uno_si_no(
    script: ModuleType, tablas: dict[str, np.ndarray], tmp_path: Path
) -> None:
    assert script.main(["--verificar"]) == 0
    alterado = {k: v.copy() for k, v in tablas.items()}
    alterado["sobrev_M_aa"][60, 0] += 0.01
    ruta = tmp_path / "alterado.npz"
    np.savez_compressed(ruta, **alterado)
    assert script.main(["--verificar", "--salida", str(ruta)]) == 1
