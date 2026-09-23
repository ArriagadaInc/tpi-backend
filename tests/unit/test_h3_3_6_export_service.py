"""Unit tests for ``SolicitudService.export_leads_xlsx`` (fake repository).

Verifies authorization fail-closed, count-first limit abort, filter forwarding,
PII-free audit metadata and the explicit allowlist (extra row keys are dropped).
"""

from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any, cast
from uuid import UUID

import pytest
from openpyxl import load_workbook

from app.auth.models import AuthenticatedUser, UserRole
from app.services.solicitud_service import SolicitudService
from app.services.xlsx_export import EXPORT_MAX_ROWS, ExportLimitExceededError

AFP_ID = UUID("b8ba2d12-2de0-41a5-8349-77cda60a14b6")
ASESOR_ID = UUID("aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa")


def _user(role: str) -> AuthenticatedUser:
    return AuthenticatedUser(
        subject=f"subject-{role}",
        username=f"{role}@example.com",
        display_name=f"User {role}",
        role=cast(UserRole, role),
    )


def _row(**overrides: Any) -> dict[str, Any]:
    row: dict[str, Any] = {
        "id_lead": "11111111-1111-1111-1111-111111111111",
        "rut": "12.345.678-5",
        "nombre_completo": "Juan Perez",
        "email": "juan@example.com",
        "telefono": "+56 9 1234 5678",
        "genero": "Masculino",
        "estado_civil": "Soltero",
        "afp": "Habitat",
        "saldo_afp": 1000000,
        "comentarios": "Solicitud original.",
        "estado_lead": "contactado",
        "created_at": datetime(2026, 9, 1, 10, 0, tzinfo=UTC),
        "id_asesor": str(ASESOR_ID),
        "asesor_nombre": "Asesor Uno",
    }
    row.update(overrides)
    return row


class _ExportRepository:
    def __init__(self, rows: list[dict[str, Any]] | None = None) -> None:
        self.rows = rows if rows is not None else [_row()]
        self.count_calls: list[dict[str, Any]] = []
        self.export_calls: list[dict[str, Any]] = []
        self.audit_calls: list[dict[str, Any]] = []
        self.count_override: int | None = None

    def count_crm_solicitudes(self, **kwargs: Any) -> int:
        self.count_calls.append(dict(kwargs))
        return self.count_override if self.count_override is not None else len(self.rows)

    def get_crm_solicitudes_export(self, **kwargs: Any) -> list[dict[str, Any]]:
        self.export_calls.append(dict(kwargs))
        return [dict(row) for row in self.rows]

    def record_xlsx_export_event(self, **kwargs: Any) -> None:
        self.audit_calls.append(dict(kwargs))


def _service(repo: _ExportRepository) -> tuple[SolicitudService, _ExportRepository]:
    return SolicitudService(repository=cast(Any, repo)), repo


def test_superuser_export_returns_xlsx_bytes() -> None:
    service, repo = _service(_ExportRepository())
    content = service.export_leads_xlsx(_user("ceo"))
    assert isinstance(content, bytes)
    assert content[:2] == b"PK"
    assert len(repo.audit_calls) == 1


@pytest.mark.parametrize(
    "role", ["advisor", "admin", "executive", "operations", "readonly", "tester"]
)
def test_non_superuser_export_raises_permission_error_before_query(role: str) -> None:
    service, repo = _service(_ExportRepository())
    with pytest.raises(PermissionError):
        service.export_leads_xlsx(_user(role))
    assert repo.count_calls == []
    assert repo.export_calls == []
    assert repo.audit_calls == []


def test_limit_exceeded_aborts_before_fetch() -> None:
    repo = _ExportRepository()
    repo.count_override = EXPORT_MAX_ROWS + 1
    service, _ = _service(repo)
    with pytest.raises(ExportLimitExceededError):
        service.export_leads_xlsx(_user("cto"))
    assert repo.count_calls != []
    assert repo.export_calls == []
    assert repo.audit_calls == []


def test_limit_exact_is_allowed_and_exports() -> None:
    # AC-16: count == EXPORT_MAX_ROWS must export (only count > limit aborts).
    repo = _ExportRepository()
    repo.count_override = EXPORT_MAX_ROWS
    service, _ = _service(repo)
    content = service.export_leads_xlsx(_user("ceo"))
    assert isinstance(content, bytes)
    assert content[:2] == b"PK"
    assert repo.export_calls != []
    assert len(repo.audit_calls) == 1


def test_filters_forwarded_to_repository() -> None:
    repo = _ExportRepository()
    service, _ = _service(repo)
    service.export_leads_xlsx(
        _user("ceo"),
        search="Juan",
        estado_lead="contactado",
        afp_id=AFP_ID,
        date_from=date(2026, 1, 1),
        date_to=date(2026, 12, 31),
        sort_by="nombre_completo",
        sort_direction="asc",
        asesor_id=ASESOR_ID,
        origen_lead="formulario_streamlit",
        fuente_actual="backoffice",
        sin_asignar=True,
        estancado=True,
    )

    count_args = repo.count_calls[0]
    assert count_args["search"] == "Juan"
    assert count_args["estado_lead"] == "contactado"
    assert count_args["afp_id"] == AFP_ID
    assert isinstance(count_args["date_from"], datetime)
    assert isinstance(count_args["date_to"], datetime)
    assert count_args["asesor_id"] == ASESOR_ID
    assert count_args["origen_lead"] == "formulario_streamlit"
    assert count_args["fuente_actual"] == "backoffice"
    assert count_args["sin_asignar"] is True
    assert count_args["estancado"] is True

    export_args = repo.export_calls[0]
    assert export_args["sort_by"] == "nombre_completo"
    assert export_args["sort_direction"] == "asc"


def test_audit_event_sanitized_without_search_text() -> None:
    repo = _ExportRepository([_row(), _row()])
    service, _ = _service(repo)
    service.export_leads_xlsx(_user("ceo"), search="12.345.678-5")

    audit = repo.audit_calls[0]
    assert audit["actor_subject"] == "subject-ceo"
    assert audit["role"] == "ceo"
    assert audit["row_count"] == 2
    filters = audit["filters_summary"]
    assert filters["search_applied"] is True
    assert "12.345.678-5" not in str(filters)
    assert "format" not in filters  # format lives at the event top level


def test_allowlist_drops_extra_row_keys() -> None:
    polluted = _row(
        raw_payload={"rut": "12.345.678-5", "secret": "SHOULD_NOT_LEAK"},
        id_persona="persona-uuid",
        password_hash="HASH_SHOULD_NOT_LEAK",
    )
    repo = _ExportRepository([polluted])
    service, _ = _service(repo)
    content = service.export_leads_xlsx(_user("ceo"))

    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    for row in sheet.iter_rows(values_only=True):
        for cell in row:
            assert "SHOULD_NOT_LEAK" not in str(cell)
            assert "HASH_SHOULD_NOT_LEAK" not in str(cell)
    # Only the 14 allowlist headers exist.
    assert sheet.max_column == 14
